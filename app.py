import sqlite3

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from io import BytesIO
from flask import send_file

from PIL import Image, UnidentifiedImageError
from pathlib import Path
from uuid import uuid4
from  werkzeug.utils import secure_filename
from flask import Flask, render_template, request

app =  Flask(__name__) #made an application in Flask

UPLOAD_FOLDER = Path(app.static_folder)/ "uploads"
UPLOAD_FOLDER.mkdir(parents=True, exist_ok=True)
app.config["MAX_CONTENT_LENGTH"] = 5 * 1024 * 1024

ALLOWED_EXTENSIONS = {"jpg","jpeg","png","Webp"}

def init_db():
    connection = sqlite3.connect("profitlens.db")

    connection.execute("""
            CREATE TABLE IF NOT EXISTS products (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL,
            product_cost REAL NOT NULL,
            total_cost REAL NOT NULL,
            selling_price REAL NOT NULL,
            profit REAL NOT NULL
        )
    """)
    columns = connection.execute(
        "PRAGMA table_info(products)"
    ).fetchall()
    column_names = [column[1] for column in columns]
    if "image_filename" not in column_names:
        connection.execute("""
            ALTER TABLE products
            ADD COLUMN image_filename TEXT
        """)
    connection.commit()
    connection.close()


def allowed_file(filename):
    return (
        "." in filename
        and filename.rsplit(".", 1)[1].lower()
        in ALLOWED_EXTENSIONS
    )
def save_product_image(file):
    if not file or not file.filename:
        return None
    safe_name = secure_filename(file.filename)
    if not allowed_file(safe_name):
        raise ValueError("Please upload a JPG, PNG, or WebP image.")

    try:
        image = Image.open(file.stream)
        image.load()
    except (UnidentifiedImageError, OSError, ValueError):
        raise ValueError("The uploaded file is not a valid image.")

    if image.format not in {"JPEG","PNG","WEBP"}:
        raise ValueError("Unsupported image format.")

    filename = f"{uuid4().hex}.png"
    destination = UPLOAD_FOLDER / filename 

    image = image.convert("RGB")
    image.save(destination, format = "PNG")

    return filename



@app.route("/")
def home():

    selected_price = request.args.get("price", "")
    selected_shipping = request.args.get("supplier_shipping", "")
    selected_link = request.args.get("supplier_link", "") 
    return render_template(
        "index.html",
        result=None,
        selected_price=selected_price,
        selected_shipping=selected_shipping,
        selected_link=selected_link
    )


@app.route("/calculate", methods=["POST"])

def calculate():
    name = request.form["supplier_link"]
    price = float(request.form["price"])
    rate = float(request.form["rate"])
    shipping = float(request.form["supplier_shipping"]or 0)
    other_costs = float(request.form["other_costs"])
    markup = float(request.form["profit"])
    product_cost = price*rate
    total_cost = (price + shipping)*rate +other_costs
    profit = round(total_cost * markup /100,2)
    selling_price = round(total_cost + profit,2)
    
    uploaded_file = request.files.get("product_image")

    try:
        image_filename = save_product_image(uploaded_file)
    except ValueError as error:
        return str(error), 400
    
    connection = sqlite3.connect("profitlens.db")

    connection.execute("""
    INSERT INTO products
    (name, product_cost, total_cost, selling_price, profit, image_filename)
    VALUES (?,?,?,?,?,?)    
    """,(
        name, 
        product_cost,
        total_cost,
        selling_price,
        profit,
        image_filename
    ))
    connection.commit()
    connection.close()

    result = {
        "name": name,
        "product_cost": round(product_cost, 2),
        "total_cost": round(total_cost, 2),
        "profit": round(profit, 2),
        "selling_price": round(selling_price, 2)
    }
    return render_template("index.html", result=result)


@app.route("/inventory")
def inventory():
    connection = sqlite3.connect("profitlens.db")

    products = connection.execute("""
    SELECT id,name,total_cost, selling_price, profit, image_filename
    FROM products    
""").fetchall()
    connection.close()
    return render_template("inventory.html", products=products)



@app.route("/compare", methods=["POST"])
def compare():
    suppliers = []
    for i in range(1,4):
        link = request.form.get(f"link{i}","").strip()
        price_text = request.form.get(f"price{i}","").strip()
        shipping_text = request.form.get(f"shipping{i}", "0").strip()

        if not link and not price_text:
            continue
        if not link or not price_text:
            return "Please provide both a link and a price.", 400

        try:
            price = float(price_text)
            shipping = float(shipping_text or "0")

        except ValueError:
            return "Please enter valid numbers.",400

        if price < 0 or shipping < 0:
            return "Prices cannot be negative.",400

        total= price + shipping

        suppliers.append({
            "link": link,
            "price": price,
            "shipping": shipping,
            "total": round(total,2)
        })
    if not suppliers:
        return "Please enter at least one supplier.",400

    cheapest = min(
        suppliers,
        key=lambda supplier: supplier["total"]
    )

    return render_template(
        "compare.html",
        suppliers=suppliers,
        cheapest=cheapest
    )


@app.route("/export")
def export():

    connection = sqlite3.connect("profitlens.db")

    products = connection.execute("""
        SELECT id, name, product_cost,
               total_cost, selling_price, profit
        FROM products
    """).fetchall()

    connection.close()

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "ProfitLens Inventory"

    headers = [
        "ID",
        "AliExpress Link",
        "Product Cost (BDT)",
        "Total Cost (BDT)",
        "Selling Price (BDT)",
        "Profit (BDT)"
    ]

    sheet.append(headers)

    for product in products:
        sheet.append(list(product))

    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="1F4E78"
        )
        cell.alignment = Alignment(horizontal="center")

    widths = [10, 40, 22, 22, 24, 20]

    for index, width in enumerate(widths, start=1):
        letter = get_column_letter(index)
        sheet.column_dimensions[letter].width = width

    for row in sheet.iter_rows(min_row=2):
        link_cell = row[1]
        if link_cell.value:
            link_cell.hyperlink = link_cell.value
            link_cell.value = "View Product"
            link_cell.style = "Hyperlink"

        for cell in row[2:]:
            cell.number_format = '#,##0.00'

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    return send_file(
        output,
        as_attachment=True,
        download_name="ProfitLens_Inventory.xlsx",
        mimetype=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )




if __name__== "__main__":
    init_db()
    app.run(debug=True)